from io import BytesIO
from zipfile import BadZipFile

import pymupdf
from docx import Document as WordDocument
from fastapi import HTTPException

from app.documents.validation import fail

MAX_PDF_PAGES = 2_000
MAX_EXTRACTED_CHARACTERS = 10_000_000


def extract_pages(filename: str, content: bytes) -> list[dict]:
    suffix = filename.rsplit(".", 1)[-1].lower()
    try:
        if suffix == "pdf":
            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                if pdf.is_encrypted:
                    raise fail(
                        422,
                        "PDF_ENCRYPTED",
                        "Password-protected PDFs are not supported",
                    )
                if pdf.page_count > MAX_PDF_PAGES:
                    raise fail(
                        422, "PDF_PAGE_LIMIT", "The PDF has too many pages to process"
                    )
                pages = [
                    {"page": number, "text": page.get_text("text", sort=True).strip()}
                    for number, page in enumerate(pdf, start=1)
                ]
        elif suffix == "docx":
            word = WordDocument(BytesIO(content))
            text = [
                paragraph.text.strip()
                for paragraph in word.paragraphs
                if paragraph.text.strip()
            ]
            for table in word.tables:
                for row in table.rows:
                    cells = [
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    ]
                    if cells:
                        text.append(" | ".join(cells))
            pages = [{"page": None, "text": "\n".join(text)}]
        elif suffix == "txt":
            pages = [
                {
                    "page": None,
                    "text": content.decode("utf-8-sig", errors="strict").strip(),
                }
            ]
        else:
            raise fail(
                415,
                "FILE_TYPE_UNSUPPORTED",
                "Only PDF, DOCX, and TXT files are supported",
            )
    except HTTPException:
        raise
    except (
        pymupdf.FileDataError,
        pymupdf.EmptyFileError,
        BadZipFile,
        ValueError,
        RuntimeError,
    ):
        raise fail(
            422, "DOCUMENT_PARSE_FAILED", "The document could not be safely parsed"
        ) from None
    total = sum(len(item["text"]) for item in pages)
    if total > MAX_EXTRACTED_CHARACTERS:
        raise fail(
            422,
            "EXTRACTED_TEXT_LIMIT",
            "The extracted document text exceeds the processing limit",
        )
    if suffix == "pdf" and not any(
        sum(ch.isalnum() for ch in item["text"]) >= 20 for item in pages
    ):
        raise fail(
            422,
            "OCR_REQUIRED",
            "This PDF has no extractable text; scanned-document OCR is not available",
        )
    if not any(item["text"] for item in pages):
        raise fail(
            422, "DOCUMENT_TEXT_EMPTY", "No readable text was found in the document"
        )
    return pages
